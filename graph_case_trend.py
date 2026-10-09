# ケース単位の推移グラフ（平均・最小〜最大の帯・はね本数）を全ロット分まとめて出力する
# 使い方: 下の「設定」を書き換えて  python graph_case_trend.py
# 必要なもの: pip install openpyxl matplotlib numpy  / 日本語フォント(Noto Sans CJK JP など)

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager as fm
from openpyxl import load_workbook

# ===== 設定（ここだけ書き換えればOK）=====
EXCEL_FILE = "Book1.xlsx"                      # 読み込むExcel
OUT_FILE   = "全ロット_ケース単位の推移.png"     # 出力するPNG
N_LOTS     = 15                                # ロット数
FIRST_ROW  = 5                                 # データが始まる行（1ケース目の行）
LAST_ROW   = 37                                # データが終わる行
FIRST_COL  = 1                                 # ロット1の「平均」列（0始まり。B列=1）
COLS_PER_LOT = 4                               # 1ロットあたりの列数（平均・最小・最大・はね本数）
FONT_NAME = "Hiragino Sans"                # 日本語フォント名（Windowsなら "MS Gothic" や "Yu Gothic"）
# ==========================================

plt.rcParams["font.family"] = FONT_NAME
plt.rcParams["axes.unicode_minus"] = False

# --- Excelから各ロットのデータを読む（数式の計算結果を読むため data_only=True）---
ws = load_workbook(EXCEL_FILE, data_only=True).active
rows = list(ws.iter_rows(min_row=FIRST_ROW, max_row=LAST_ROW, values_only=True))
lots = []
for k in range(N_LOTS):
    c = FIRST_COL + COLS_PER_LOT * k
    # 空欄のケース行は除く → 列は [平均, 最小, 最大, はね本数]
    d = np.array([r[c:c + 4] for r in rows if r[c] is not None], float)
    lots.append(d)

# --- 図の枠を用意（3列 × 必要な行数）---
COLS = 3
NR = int(np.ceil(N_LOTS / COLS))
H = 4.15 * NR + 2.0
fig = plt.figure(figsize=(18, H), dpi=130)
gs = fig.add_gridspec(NR, COLS, left=0.045, right=0.955, top=1 - 1.85 / H,
                      bottom=0.035, hspace=0.62, wspace=0.30)

RIGHT_MAX = 95   # 右軸（はね本数）の上限。全ロット共通にして見比べやすくする

for i, d in enumerate(lots):
    ax = fig.add_subplot(gs[i // COLS, i % COLS])
    x = np.arange(1, len(d) + 1)                 # ケース番号 1,2,3...
    mean, mn, mx, rej = d[:, 0], d[:, 1], d[:, 2], d[:, 3]

    # 水色の帯＝ケース内の最小〜最大、濃い線＝ケースごとの平均
    ax.fill_between(x, mn, mx, color="#90caf9", alpha=.6, lw=0, label="そのケース内の最小〜最大")
    ax.plot(x, mean, "-o", color="#0d47a1", lw=2.6, ms=5.5, label="ケースごとの平均", zorder=3)

    # 縦軸の範囲：下に余白を作って、はね本数の棒が帯と重なりにくいようにする
    lo, hi = mn.min(), mx.max()
    span = hi - lo
    ax.set_ylim(lo - 0.75 * span, hi + 0.3 * span)
    ax.set_xlim(0, max(len(d) + 1.5, 31.5))
    ax.set_xticks([5, 10, 15, 20, 25, 30])
    ax.yaxis.set_major_locator(plt.MaxNLocator(5, integer=True))
    ax.grid(axis="y", color="#e8e8e8", lw=.8)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    if i % COLS == 0:
        ax.set_ylabel("製品の重さ (g)", fontsize=11)
    if i + COLS >= N_LOTS:                        # 一番下の段だけ横軸ラベル
        ax.set_xlabel("瓶ケースの順番（1ケース=80本）", fontsize=12)

    # 赤い棒＝はね本数（右軸）
    ax2 = ax.twinx()
    ax2.bar(x, rej, color="#e57373", alpha=.45, width=.7, zorder=1)
    ax2.set_ylim(0, RIGHT_MAX)
    ax2.set_yticks([0, 20, 40])
    ax2.tick_params(axis="y", colors="#c62828")
    for s in ("top", "left"):
        ax2.spines[s].set_visible(False)
    ax2.spines["right"].set_color("#c62828")
    ax.set_zorder(ax2.get_zorder() + 1)           # 線を棒の手前に出す
    ax.patch.set_visible(False)

    # タイトル：ロット名＋はね本数の合計
    ax.text(0, 1.06, f"ロット{i + 1}　はね {int(rej.sum())}本",
            transform=ax.transAxes, fontsize=16, fontweight="bold", va="bottom")
    if i == 0:
        ax.legend(loc="lower left", fontsize=10, framealpha=.95)

fig.suptitle(f"{N_LOTS}ロットのケース単位の推移", fontsize=24, fontweight="bold", y=0.993)
fig.text(0.5, 1 - 0.95 / H, "濃い線＝ケースごとの平均　／　水色の帯＝そのケースの中の最小〜最大　／　赤い棒＝はねた本数（右目盛り）",
         ha="center", fontsize=12.5, color="#555")
fig.savefig(OUT_FILE, dpi=130)
print("保存しました:", OUT_FILE)
