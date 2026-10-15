# 1本ごとの時系列グラフ（実測値・基準・上限・下限）を全ロット分まとめて出力する
# 使い方: 下の「設定」を書き換えて  python graph_timeseries.py
# 必要なもの: pip install openpyxl matplotlib numpy  / 日本語フォント(Noto Sans CJK JP など)

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager as fm
from openpyxl import load_workbook

# ===== 設定（ここだけ書き換えればOK）=====
EXCEL_FILE = "新規_Microsoft_Excel_ワークシート.xlsx"
OUT_FILE   = "全ロット_1本ごとの時系列.png"
N_LOTS     = 15      # ロット数
FIRST_ROW  = 5       # データが始まる行
FIRST_COL  = 1       # ロット1の「実測値」列（0始まり。B列=1）
COLS_PER_LOT = 4     # 1ロットあたりの列数（実測値・基準・上限・下限）
FONT_NAME = "Yu Gothic"    # Windowsなら "MS Gothic" や "Yu Gothic"
# ==========================================

plt.rcParams["font.family"] = FONT_NAME
plt.rcParams["axes.unicode_minus"] = False

ws = load_workbook(EXCEL_FILE, data_only=True).active
rows = list(ws.iter_rows(min_row=FIRST_ROW, values_only=True))

def is_num(v):
    return isinstance(v, (int, float))

COLS = 3
NR = int(np.ceil(N_LOTS / COLS))
H = 3.9 * NR + 1.8
fig = plt.figure(figsize=(18, H), dpi=130)
gs = fig.add_gridspec(NR, COLS, left=0.045, right=0.915, top=1 - 1.7 / H,
                      bottom=0.04, hspace=0.40, wspace=0.42)

for k in range(N_LOTS):
    c = FIRST_COL + COLS_PER_LOT * k
    # 4列とも数値の行だけ使う（#N/A や空欄の行は飛ばす）
    d = np.array([r[c:c + 4] for r in rows if all(is_num(r[c + j]) for j in range(4))], float)
    y = d[:, 0]
    base, upper, lower = d[0, 1], d[0, 2], d[0, 3]   # 基準・上限・下限は先頭行の値を使う
    n = len(y)

    ax = fig.add_subplot(gs[k // COLS, k % COLS])
    ax.plot(np.arange(1, n + 1), y, color="#78909c", lw=.6)            # 灰色＝1本ごとの実測値
    ax.axhline(base, color="#1565c0", lw=2.2)                           # 青＝基準
    ax.axhline(upper, color="#d32f2f", lw=2, ls="--")                   # 赤破線＝上限
    ax.axhline(lower, color="#d32f2f", lw=2, ls="--")                   # 赤破線＝下限
    # 右側に「上限 295.8」などの文字を出す
    for v, t, col in ((upper, f"上限 {upper:.1f}", "#d32f2f"),
                      (base,  f"基準 {base:.1f}",  "#1565c0"),
                      (lower, f"下限 {lower:.1f}", "#d32f2f")):
        ax.text(1.0, v, t, transform=ax.get_yaxis_transform(), color=col,
                fontsize=11, va="center", ha="left", clip_on=False)

    lo, hi = min(y.min(), lower), max(y.max(), upper)
    ax.set_ylim(lo - 0.15 * (hi - lo) - 0.5, hi + 0.15 * (hi - lo) + 0.5)
    ax.set_xlim(0, n + 20)
    ax.yaxis.set_major_locator(plt.MaxNLocator(5, integer=True))
    ax.grid(axis="y", color="#ececec", lw=.8)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_title(f"ロット{k + 1}", loc="left", fontsize=17, fontweight="bold", pad=8)
    if k % COLS == 0:
        ax.set_ylabel("製品の重さ (g)", fontsize=11)
    if k + COLS >= N_LOTS:
        ax.set_xlabel("測定順（本）", fontsize=12)

fig.suptitle("各ロットの時系列（基準値＝開始時5個の平均、上下限 ±1.4 g）",
             fontsize=24, fontweight="bold", y=1 - 0.25 / H)
fig.text(0.5, 1 - 0.85 / H, "灰色の線＝1本ごとの実測値　／　青い線＝基準値　／　赤い破線＝上限・下限",
         ha="center", fontsize=13, color="#555")
fig.savefig(OUT_FILE, dpi=130)
print("保存しました:", OUT_FILE)
